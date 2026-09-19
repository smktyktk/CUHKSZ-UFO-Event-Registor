from pathlib import Path
from datetime import datetime
import shutil
import sys

from openpyxl import load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

SIGNUP_FOLDER = BASE_DIR / "signup_sheet"
HISTORY_FOLDER = BASE_DIR / "record_history"

RECORD_FILE = BASE_DIR / "record.xlsx"
LOG_FILE = BASE_DIR / "log.txt"


# ============================================================
# BASIC UTILITIES
# ============================================================

def log(message=""):
    """
    Print a message to the console and append it to log.txt.

    Every output line receives its own timestamp in the format:
        [yyyy/mm/dd hh:mm:ss]: log content
    """

    timestamp = datetime.now().strftime(
        "%Y/%m/%d %H:%M:%S"
    )

    # Handle multi-line messages line by line so every output line
    # receives its own timestamp.
    lines = str(message).splitlines()

    if not lines:
        lines = [""]

    with LOG_FILE.open(
        "a",
        encoding="utf-8"
    ) as log_file:

        for line in lines:

            formatted = (
                f"[{timestamp}]: {line}"
            )

            print(formatted)
            log_file.write(
                formatted + "\n"
            )

def error(message):

    log()
    log("=" * 60)
    log("ERROR")
    log("=" * 60)
    log(message)
    log()
    sys.exit(1)


def normalize_student_id(value):

    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


# ============================================================
# DIRECTORY SETUP
# ============================================================

def create_directories():

    SIGNUP_FOLDER.mkdir(
        parents=True,
        exist_ok=True
    )

    HISTORY_FOLDER.mkdir(
        parents=True,
        exist_ok=True
    )


# ============================================================
# FILENAME PARSING
# ============================================================

def parse_signup_filename(file):

    prefix = "signup_sheet_"

    if not file.stem.startswith(prefix):
        return None

    remainder = file.stem[
        len(prefix):
    ]

    parts = remainder.split("_")

    if len(parts) != 2:
        return None

    date_string = parts[0]
    serial_string = parts[1]

    if len(date_string) != 8:
        return None

    try:

        serial = int(
            serial_string
        )

    except ValueError:

        return None

    return (
        date_string,
        serial
    )


# ============================================================
# FIND SIGNUP SHEETS
# ============================================================

def get_signup_sheets():

    sheets = []

    for file in SIGNUP_FOLDER.glob(
        "signup_sheet_*.xlsx"
    ):

        parsed = parse_signup_filename(
            file
        )

        if parsed is not None:

            sheets.append(
                (
                    parsed,
                    file
                )
            )

    # Sort by date and serial number.
    sheets.sort(
        key=lambda x: x[0]
    )

    return sheets


# ============================================================
# LOAD ALL ATTENDANCE DATA
# ============================================================

def calculate_records():

    sheets = get_signup_sheets()

    if not sheets:

        log()
        log(
            "No signup sheets were found."
        )

        return {}

    attendance = {}

    log()
    log(
        f"Found {len(sheets)} signup sheet(s)."
    )

    log()

    # --------------------------------------------------------
    # Process each event
    # --------------------------------------------------------

    for (date_serial, file) in sheets:

        date_string, serial = date_serial

        log(
            f"Processing "
            f"{file.name}..."
        )

        try:

            workbook = load_workbook(
                file,
                data_only=True
            )

        except Exception as e:

            error(
                f"Could not open:\n"
                f"  {file}\n\n"
                f"{e}"
            )

        worksheet = workbook.active

        # Require at least 3 columns.
        if worksheet.max_column < 3:

            workbook.close()

            error(
                f"{file.name} does not have "
                f"the required three columns:\n\n"
                "Column 1 = nickname\n"
                "Column 2 = student ID\n"
                "Column 3 = attendance"
            )

        event_student_count = 0
        event_absent_count = 0

        # ----------------------------------------------------
        # Process students
        # ----------------------------------------------------

        for row_number in range(
            2,
            worksheet.max_row + 1
        ):

            nickname = worksheet.cell(
                row=row_number,
                column=1
            ).value

            student_id = normalize_student_id(
                worksheet.cell(
                    row=row_number,
                    column=2
                ).value
            )

            mark = worksheet.cell(
                row=row_number,
                column=3
            ).value

            # Ignore completely empty rows.
            if not student_id:
                continue

            event_student_count += 1

            # ------------------------------------------------
            # Attendance rule:
            #
            # blank -> attended
            # 1     -> absent
            # ------------------------------------------------

            is_absent = False

            if mark is not None:

                # Accept numeric 1
                if mark == 1:
                    is_absent = True

                # Also accept text "1"
                elif str(mark).strip() == "1":
                    is_absent = True

            if is_absent:

                event_absent_count += 1

            else:

                # Student attended.
                attendance[student_id] = (
                    attendance.get(
                        student_id,
                        0
                    ) + 1
                )

        workbook.close()

        event_attended_count = (
            event_student_count
            - event_absent_count
        )

        log(
            f"  Students: {event_student_count}"
        )

        log(
            f"  Attended: {event_attended_count}"
        )

        log(
            f"  Absent:   {event_absent_count}"
        )

    return attendance


# ============================================================
# BACKUP OLD RECORD
# ============================================================

def get_next_history_serial(
    date_string
):

    existing_numbers = []

    pattern = (
        f"record_history_"
        f"{date_string}_*.xlsx"
    )

    for file in HISTORY_FOLDER.glob(
        pattern
    ):

        stem = file.stem

        try:

            number = int(
                stem.split("_")[-1]
            )

            if number >= 1:
                existing_numbers.append(
                    number
                )

        except ValueError:
            pass

    if not existing_numbers:
        return 1

    return max(existing_numbers) + 1


def backup_record():

    if not RECORD_FILE.exists():

        error(
            "record.xlsx does not exist."
        )

    today = datetime.now()

    date_string = today.strftime(
        "%Y%m%d"
    )

    serial = get_next_history_serial(
        date_string
    )

    history_file = (
        HISTORY_FOLDER
        / f"record_history_"
        f"{date_string}_"
        f"{serial:03d}.xlsx"
    )

    try:

        shutil.copy2(
            RECORD_FILE,
            history_file
        )

    except OSError as e:

        error(
            "Could not back up record.xlsx.\n\n"
            f"{e}"
        )

    return history_file


# ============================================================
# UPDATE RECORD
# ============================================================

def write_new_record(
    attendance
):

    if not RECORD_FILE.exists():

        error(
            "record.xlsx does not exist."
        )

    workbook = load_workbook(
        RECORD_FILE
    )

    worksheet = workbook.active

    # --------------------------------------------------------
    # Preserve the header.
    # --------------------------------------------------------

    # Remove all existing data rows.
    if worksheet.max_row > 1:

        worksheet.delete_rows(
            2,
            worksheet.max_row - 1
        )

    # --------------------------------------------------------
    # Sort student IDs
    # --------------------------------------------------------

    # Using string sorting keeps IDs such as 00123 intact.
    sorted_students = sorted(
        attendance.keys()
    )

    # --------------------------------------------------------
    # Write new record
    # --------------------------------------------------------

    for row_number, student_id in enumerate(
        sorted_students,
        start=2
    ):

        worksheet.cell(
            row=row_number,
            column=1
        ).value = student_id

        worksheet.cell(
            row=row_number,
            column=2
        ).value = attendance[
            student_id
        ]

    try:

        workbook.save(
            RECORD_FILE
        )

    except Exception as e:

        workbook.close()

        error(
            "Could not save the updated record.xlsx.\n\n"
            f"{e}"
        )

    workbook.close()


# ============================================================
# MAIN
# ============================================================

def main():

    log()
    log("=" * 60)
    log("WJX ATTENDANCE RECORD PROCESSOR")
    log("=" * 60)

    # --------------------------------------------------------
    # Setup
    # --------------------------------------------------------

    create_directories()

    # --------------------------------------------------------
    # Check record.xlsx
    # --------------------------------------------------------

    if not RECORD_FILE.exists():

        error(
            f"record.xlsx was not found.\n\n"
            f"Expected location:\n"
            f"  {RECORD_FILE}"
        )

    # --------------------------------------------------------
    # Find signup sheets
    # --------------------------------------------------------

    sheets = get_signup_sheets()

    if not sheets:

        log()
        log(
            "No signup sheets were found."
        )

        log()
        log(
            "record.xlsx has not been modified."
        )

        return

    # --------------------------------------------------------
    # Display sheets
    # --------------------------------------------------------

    log()
    log(
        f"Found {len(sheets)} signup sheet(s):"
    )

    for date_serial, file in sheets:

        log(
            f"  {file.name}"
        )

    # --------------------------------------------------------
    # Calculate attendance from scratch
    # --------------------------------------------------------

    log()
    log(
        "Recalculating attendance "
        "from all signup sheets..."
    )

    attendance = calculate_records()

    # --------------------------------------------------------
    # Backup current record
    # --------------------------------------------------------

    log()
    log(
        "Backing up current record.xlsx..."
    )

    history_file = backup_record()

    log(
        f"Backup created:\n"
        f"  {history_file}"
    )

    # --------------------------------------------------------
    # Write new record
    # --------------------------------------------------------

    log()
    log(
        "Writing recalculated record.xlsx..."
    )

    write_new_record(
        attendance
    )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    log()
    log("=" * 60)
    log("RECORD UPDATE COMPLETED")
    log("=" * 60)

    log()
    log(
        f"Students with at least one attendance: "
        f"{len(attendance)}"
    )

    log()
    log(
        f"Updated record:\n"
        f"  {RECORD_FILE}"
    )

    log()
    log(
        f"History backup:\n"
        f"  {history_file}"
    )

    log()
    log("=" * 60)


if __name__ == "__main__":
    main()